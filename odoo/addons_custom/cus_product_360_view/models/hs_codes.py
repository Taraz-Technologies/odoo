from odoo import _, api, fields, models
from odoo.exceptions import UserError

import PyPDF2
from openpyxl import load_workbook, Workbook


def isfloat(num):
    try:
        float(num)
        return True
    except ValueError:
        return False


class CountryScenarioRules(models.Model):
    _name = "country.scenario.rule"
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _description = "Country Scenario Rules"
    _rec_name = "display_name"

    x_label = fields.Char(string='Label', required=False)
    x_country_id = fields.Many2one(comodel_name='res.country', string='Country', tracking=True)
    x_currency_id = fields.Many2one(comodel_name='res.currency', string='Currency', default=2, tracking=True)
    x_default_scenario = fields.Boolean(string='Default', required=False)
    x_landed_cost = fields.Float(string='Landing Cost(%)', tracking=True)
    x_insurance = fields.Float(string='Insurance(%)', tracking=True)

    x_cd_formula = fields.Char(string='CD Formula', tracking=True)
    x_acd_formula = fields.Char(string='ACD Formula', tracking=True)
    x_rd_formula = fields.Char(string='RD Formula', tracking=True)
    x_st_formula = fields.Char(string='ST Formula', tracking=True)
    x_ast_formula = fields.Char(string='AST Formula', tracking=True)
    x_it_formula = fields.Char(string='IT Formula', tracking=True)
    x_cess_formula = fields.Char(string='CESS Formula', tracking=True)

    x_av_label = fields.Char(string='AV Label', tracking=True)
    x_cd_label = fields.Char(string='CD Label', tracking=True)
    x_acd_label = fields.Char(string='ACD Label', tracking=True)
    x_rd_label = fields.Char(string='RD Label', tracking=True)
    x_st_label = fields.Char(string='ST Label', tracking=True)
    x_ast_label = fields.Char(string='AST Label', tracking=True)
    x_it_label = fields.Char(string='IT Label', tracking=True)
    x_cess_label = fields.Char(string='CESS Label', tracking=True)

    x_manager_salary = fields.Float(string='Manager Salary', tracking=True)
    x_engineer_salary = fields.Float(string='Engineer Salary', tracking=True)
    x_technician_salary = fields.Float(string='Technician Salary', tracking=True)

    x_per_square_rent = fields.Float(string='Per Square Rent', tracking=True)
    x_per_kw_rate = fields.Float(string='Per kW Rate', tracking=True)
    x_miscellaneous = fields.Float(string='Misc.', tracking=True)

    x_notes = fields.Text(string="Notes", required=False)

    display_name = fields.Char(compute="_compute_display_name", store=True)

    @api.depends('x_label', 'x_default_scenario')
    def _compute_display_name(self):
        for rec in self:
            if rec.x_default_scenario:
                rec.display_name = '%s (%s)' % (rec.x_label, 'Default')
            else:
                rec.display_name = rec.x_label

    @api.onchange('x_default_scenario')
    def update_default_scenario(self):
        for rec in self:
            if rec.x_default_scenario:
                scenario_ids = self.env['country.scenario.rule'].search([])
                for scenario in scenario_ids:
                    scenario.x_default_scenario = False
                rec.x_default_scenario = True
                # product_ids = self.env['product.template'].search([('categ_id', 'in', [269, 271])])
                # for product in product_ids:
                #     if product.dk_htsus_code:
                #         hs_code_ids = product.dk_htsus_code.x_hs_code_ids
                #         product.dk_hs_code = hs_code_ids.filtered(lambda l: l.x_country_id.id == rec.x_country_id.id)
            else:
                raise UserError('Error: Make any other scenario default to change default scenario!')


class CountryImportZones(models.Model):
    _name = "country.import.zones"
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _description = "Country Zones"
    _rec_name = "display_name"

    x_country_id = fields.Many2one(comodel_name='res.country', string='Country', required=True, tracking=True)
    x_name = fields.Char(string='Zone Name', required=True, tracking=True)
    x_type = fields.Selection(selection=[
        ('import', 'Import'),
        ('export', 'Export'), ], string='Zone Type', required=True, tracking=True, default='import')

    x_country_ids = fields.Many2many('res.country', 'res_country_country_import_zones_rel', 'res_country_id',
                                     'res_country_import_zones_id', string='Zone Countries', tracking=True)

    display_name = fields.Char(compute="_compute_display_name", store=True)

    @api.depends('x_name', 'x_country_id')
    def _compute_display_name(self):
        for rec in self:
            rec.display_name = '%s (%s)' % (rec.x_name, rec.x_country_id.code)


class HsCodes(models.Model):
    _name = "hs.codes"
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _description = "HTSUS Code Database"
    _rec_name = "x_htsus_code_id"

    x_htsus_code_id = fields.Many2one(comodel_name='hs.code.database', string='HTSUS Code',
                                      required=True, domain="[('x_country_id', '=', 233)]")
    x_uom_id = fields.Many2one(comodel_name="uom.uom", string="UoM", related="x_htsus_code_id.x_uom_id")

    x_exclude = fields.Boolean(string='Not Relevant?', tracking=True)
    x_mark_red = fields.Boolean(string='Mark Red?', tracking=True)

    x_alternate_htsus = fields.Char(string="Alternate HTSUS", )

    x_tag_ids = fields.Many2many(comodel_name='hs.code.tag', string='Tags')

    x_htsus_code_group = fields.Char(string="HTSUS Code Group", compute="_compute_htsus_code_group", store=True)
    x_htsus_code_simple = fields.Char(string="HTSUS Code Simple", compute="_compute_htsus_code_group", store=True)
    x_chapter_description = fields.Char(string="Chapter Description", related="x_htsus_code_id.x_chapter_description", )
    x_heading_description = fields.Char(string="Heading", related="x_htsus_code_id.x_heading_description", )
    x_level_1_description = fields.Char(string="Sub Heading 1", related="x_htsus_code_id.x_level_1_description", )
    x_level_2_description = fields.Char(string="Sub Heading 2", related="x_htsus_code_id.x_level_2_description", )
    x_level_3_description = fields.Char(string="Sub Heading 3", related="x_htsus_code_id.x_level_3_description", )
    x_level_4_description = fields.Char(string="Sub Heading 4", related="x_htsus_code_id.x_level_4_description", )
    x_level_5_description = fields.Char(string="Sub Heading 5", related="x_htsus_code_id.x_level_5_description", )
    x_level_6_description = fields.Char(string="Sub Heading 6", related="x_htsus_code_id.x_level_6_description", )
    x_hs_code_description = fields.Char(string="HS Code Desc.", related="x_htsus_code_id.x_hs_code_description", )

    x_hs_code_ids = fields.Many2many(comodel_name='hs.code.database', string='HS Codes by Country', tracking=True)
    x_product_ids = fields.Many2many(comodel_name='product.template', string='Products', tracking=True,
                                     compute="update_product_hs_codes", store=True)

    x_search_unassigned = fields.Char(string='Search Unassigned', required=False)
    x_unassigned_categ_id = fields.Many2one(comodel_name='product.category', string='Category Unassigned', default=269,
                                            domain="[('id', 'in', [269, 271])]")

    x_unassigned_product_ids = fields.Many2many('product.template', 'product_template_hs_codes_rel_1',
                                                'product_template_id_1', 'hs_codes_id_1', string='Unassigned Products')

    @api.depends('x_htsus_code_id')
    def _compute_htsus_code_group(self):
        for rec in self:
            if rec.x_htsus_code_id:
                rec.x_htsus_code_group = rec.x_htsus_code_id.x_hs_code.replace('.', '')[:6]
                rec.x_htsus_code_simple = rec.x_htsus_code_id.x_hs_code.replace('.', '')

    @api.depends('x_hs_code_ids')
    def update_product_hs_codes(self):
        for rec in self:
            for product in rec.x_product_ids:
                if not product.dk_hs_code:
                    for hs_code in rec.x_hs_code_ids:
                        if hs_code.x_country_id.name == 'Pakistan':
                            product.dk_hs_code = hs_code.id
                            break
                if not product.x_hs_code_tr_id:
                    for hs_code in rec.x_hs_code_ids:
                        if hs_code.x_country_id.name == 'Turkey':
                            product.x_hs_code_tr_id = hs_code.id
                            break

                if product.x_htsus_assigning == 'auto':
                    if product.dk_hs_code.id not in rec.x_hs_code_ids.ids:
                        product.dk_hs_code = False
                    if product.x_hs_code_tr_id.id not in rec.x_hs_code_ids.ids:
                        product.x_hs_code_tr_id = False

    @api.onchange('x_product_ids')
    def update_hs_codes_list(self):
        for rec in self:
            h_hs_code_ids = rec.x_hs_code_ids.filtered(lambda l: l.x_country_id.name in ['Pakistan', 'Turkey']).ids

            p_hs_code_ids = rec.x_product_ids.mapped('dk_hs_code').ids
            p_hs_code_ids += rec.x_product_ids.mapped('x_hs_code_tr_id').ids

            for hs_code in h_hs_code_ids:
                if hs_code not in p_hs_code_ids:
                    rec.x_hs_code_ids = [(3, hs_code)]

            for hs_code in p_hs_code_ids:
                if hs_code not in h_hs_code_ids:
                    rec.x_hs_code_ids = [(4, hs_code)]

    def select_all(self):
        for rec in self:
            for product in rec.x_unassigned_product_ids:
                product.x_htsus_select = True

    def unselect_all(self):
        for rec in self:
            for product in rec.x_unassigned_product_ids:
                product.x_htsus_select = False

    def update_assigned(self):
        for rec in self:
            for product in rec.x_unassigned_product_ids:
                if product.x_htsus_select:
                    product.x_htsus_assigning = 'manual'
                    product.dk_htsus_code = rec.id
                    product.x_htsus_select = False
                    if not product.dk_hs_code:
                        for hs_code in rec.x_hs_code_ids:
                            if hs_code.x_country_id.name == 'Pakistan':
                                product.dk_hs_code = hs_code.id
                                break
                    if not product.x_hs_code_tr_id:
                        for hs_code in rec.x_hs_code_ids:
                            if hs_code.x_country_id.name == 'Turkey':
                                product.x_hs_code_tr_id = hs_code.id
                                break
            rec.filter_unassigned()

    @api.onchange('x_unassigned_categ_id', 'x_search_unassigned', 'x_product_ids')
    def filter_unassigned(self, ):
        for rec in self:
            domain = [('dk_htsus_code', '=', False)]
            if rec.x_unassigned_categ_id:
                domain.append(('categ_id', '=', rec.x_unassigned_categ_id.id))
            else:
                domain.append(('categ_id', 'in', [269, 271]))

            if rec.x_search_unassigned:
                domain.append('|')
                domain.append('|')
                domain.append('|')
                domain.append('|')
                domain.append(('name', 'ilike', rec.x_search_unassigned))
                domain.append(('x_taraz_part_number_id', 'ilike', rec.x_search_unassigned))
                domain.append(('description', 'ilike', rec.x_search_unassigned))
                domain.append(('dk_category', 'ilike', rec.x_search_unassigned))
                domain.append(('dk_sub_category', 'ilike', rec.x_search_unassigned))

            product_ids = self.env['product.template'].search(domain).ids
            rec.x_unassigned_product_ids = [(6, 0, product_ids)]

    def write(self, vals):
        old_hs_codes = []
        for rec in self:
            for product in rec.x_product_ids:
                if product.dk_hs_code:
                    product.dk_hs_code.x_product_ids = [(3, product.id)]
                if product.x_hs_code_tr_id:
                    product.x_hs_code_tr_id.x_product_ids = [(3, product.id)]
                if rec.x_htsus_code_id:
                    rec.x_htsus_code_id.x_product_ids = [(3, product.id)]

            old_hs_codes = rec.x_hs_code_ids

        res = super(HsCodes, self).write(vals)

        for rec in self:
            for product in rec.x_product_ids:
                if product.dk_hs_code:
                    product.dk_hs_code.x_product_ids = [(4, product.id)]
                if product.x_hs_code_tr_id:
                    product.x_hs_code_tr_id.x_product_ids = [(4, product.id)]
                if rec.x_htsus_code_id:
                    rec.x_htsus_code_id.x_product_ids = [(4, product.id)]

            new_hs_codes = rec.x_hs_code_ids

            if old_hs_codes != new_hs_codes:
                body = "<strong>HS Code by Country<strong>"

                hs_codes = old_hs_codes - new_hs_codes
                for code in hs_codes:
                    body += "<span style='color:red;font-weight:normal;'><br/>Removed: " + code.display_name + "</span>"

                hs_codes = new_hs_codes - old_hs_codes
                for code in hs_codes:
                    body += "<span style='color:green;font-weight:normal;'><br/>Added: " + code.display_name + "</span>"

                rec.message_post(message_type="comment", body=body)

        return res

    def send_record_change(self, record_name, field_names, old_values, new_values):
        for rec in self:
            body = "<strong>" + record_name + "<strong> <br/>"

            values = zip(field_names, old_values, new_values)

            for value in values:
                if value[1] != value[2]:
                    body += "<span style='color:#A00;font-weight: normal;'>&emsp;&#8226; " + value[0] + ": "
                    if not value[1] and value[2]:
                        body += value[2] + '</span> <br/>'
                    elif value[1] and not value[2]:
                        body += value[1] + ' &#8594</span> <br/>'
                    else:
                        body += value[1] + ' &#8594 ' + value[2] + '</span> <br/>'
            rec.message_post(message_type="comment", body=body)


class HsCodeDatabase(models.Model):
    _name = "hs.code.database"
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _description = "HS Code Database"
    _rec_name = 'display_name'

    x_country_id = fields.Many2one(comodel_name='res.country', string='Country', required=True, tracking=True)
    x_hs_code = fields.Char(string="HS Code", required=True, tracking=True)

    x_uom_id = fields.Many2one(comodel_name="uom.uom", string="UoM", tracking=True, )
    x_hs_code_format = fields.Char(string="HS Code", compute="_compute_hs_code", store=True)

    x_hs_code_group = fields.Char(string="Group", tracking=True, )
    x_hs_code_head = fields.Char(string="Head", tracking=True, )

    x_exclude = fields.Boolean(string='Not Relevant?', tracking=True)
    x_mark_red = fields.Boolean(string='Mark Red?', tracking=True)

    x_tag_ids = fields.Many2many(comodel_name='hs.code.tag', string='Tags')

    x_hs_code_simple = fields.Char(string="HS Code Simple", compute="_compute_hs_code", store=True)
    x_chapter_description = fields.Char(string="Chapter Desc.", tracking=True, )
    x_heading_description = fields.Char(string="Heading", tracking=True, )
    x_level_1_description = fields.Char(string="Sub Heading 1", tracking=True, )
    x_level_2_description = fields.Char(string="Sub Heading 2", tracking=True, )
    x_level_3_description = fields.Char(string="Sub Heading 3", tracking=True, )
    x_level_4_description = fields.Char(string="Sub Heading 4", tracking=True, )
    x_level_5_description = fields.Char(string="Sub Heading 5", tracking=True, )
    x_level_6_description = fields.Char(string="Sub Heading 6", tracking=True, )
    x_hs_code_description = fields.Char(string="HS Code Desc.", tracking=True, )

    x_country_tariff_ids = fields.One2many('country.tariff', 'x_hs_code_id', 'Country Tariff', tracking=True)

    x_financial_changes_ids = fields.One2many(comodel_name='financial.changes', inverse_name='x_hs_code_id',
                                              string='Financial Changes', tracking=True)
    x_product_ids = fields.Many2many(comodel_name='product.template', string='Products', tracking=True)

    display_name = fields.Char(compute="_compute_display_name", store=True)

    @api.depends('x_hs_code')
    def _compute_hs_code(self):
        for rec in self:
            if rec.x_hs_code:
                rec.x_hs_code_simple = rec.x_hs_code.replace('.', '')
                if rec.x_country_id.name == 'India':
                    rec.x_hs_code_format = rec.x_hs_code[:4] + '.' + rec.x_hs_code[4:6] + '.' + rec.x_hs_code[-2:]
                elif rec.x_country_id.name == 'Pakistan':
                    rec.x_hs_code_format = rec.x_hs_code[:7] + '.' + rec.x_hs_code[-2:]
                elif rec.x_country_id.name == 'United States':
                    rec.x_hs_code_format = rec.x_hs_code[:10] + '.' + rec.x_hs_code[-2:]
                else:
                    rec.x_hs_code_format = rec.x_hs_code

    @api.depends('x_hs_code', 'x_country_id')
    def _compute_display_name(self):
        for rec in self:
            rec.display_name = '%s (%s)' % (rec.x_hs_code, rec.x_country_id.code)

    @api.model
    def extract_hs_code_details_from_pdf(self):
        # GET HS Codes from PDF
        hs_code_data = []
        heading_description = ''
        file = open('/home/taraz/Downloads/2021101313103646911PakistanCustomsTariff-Ch1-97.pdf', 'rb')
        reader = PyPDF2.PdfFileReader(file)
        for page in range(reader.numPages):
            page = reader.pages[page]
            page_data = page.extractText()

            lines = page_data.split('\n')
            line_count = 0
            for line in lines:
                if len(line) == 5 and line[2] == '.':
                    heading_description = lines[line_count + 1]
                if len(line) == 9 and line[4] == '.':
                    hs_code_level = lines[line_count + 1].count('- ')
                    level_1_description = ''
                    if hs_code_level > 1:
                        level_line_count = line_count
                        for count in range(line_count):
                            level_count = lines[level_line_count].count('- ')
                            if level_count == 1:
                                level_1_description = lines[level_line_count]
                                break
                            level_line_count -= 1
                    level_2_description = ''
                    if hs_code_level > 2:
                        level_line_count = line_count
                        for count in range(line_count):
                            level_count = lines[level_line_count].count('- ')
                            if level_count == 2:
                                level_2_description = lines[level_line_count]
                                break
                            level_line_count -= 1
                    level_3_description = ''
                    if hs_code_level > 3:
                        level_line_count = line_count
                        for count in range(line_count):
                            level_count = lines[level_line_count].count('- ')
                            if level_count == 3:
                                level_2_description = lines[level_line_count]
                                break
                            level_line_count -= 1

                    cd_rate = 0
                    if len(lines) > line_count + 2:
                        if isfloat(lines[line_count + 2]):
                            cd_rate = lines[line_count + 2]
                        elif len(lines) > line_count + 3:
                            if isfloat(lines[line_count + 3]):
                                cd_rate = lines[line_count + 3]
                            elif len(lines) > line_count + 4:
                                if isfloat(lines[line_count + 4]):
                                    cd_rate = lines[line_count + 4]
                                elif len(lines) > line_count + 5:
                                    if isfloat(lines[line_count + 5]):
                                        cd_rate = lines[line_count + 5]

                    hs_code_data.append({
                        'hs_code': lines[line_count],
                        'heading_description': heading_description,
                        'level_1_description': level_1_description,
                        'level_2_description': level_2_description,
                        'level_3_description': level_3_description,
                        'hs_code_description': lines[line_count + 1],
                        'cd_rate': cd_rate,
                    })
                line_count += 1

        for hs_code in hs_code_data:
            hs_code_id = self.search([('x_hs_code', '=', hs_code.get('hs_code'))])
            if not hs_code_id:
                self.env['hs.code.database'].create({
                    'x_hs_code': hs_code.get('hs_code'),
                    'x_heading_description': hs_code.get('heading_description'),
                    'x_level_1_description': hs_code.get('level_1_description'),
                    'x_level_2_description': hs_code.get('level_2_description'),
                    'x_level_3_description': hs_code.get('level_3_description'),
                    'x_hs_code_description': hs_code.get('hs_code_description'),
                    'x_custom_duty': hs_code.get('cd_rate'),
                })

    @api.model
    def extract_hs_code_details_from_xl(self):
        wb_data = load_workbook('F:/Python/Excel_Work/htsdata.xlsx')
        wb_target = Workbook()

        ws_data = wb_data.active
        ws_target = wb_target.active

        hscode = None
        head = None
        chapter_desc = None
        heading = None
        sub_rows = "EFGHIJK"

        hscode_index = 0
        target_index = 2
        head_index = 0

        hscodes = []

        for i, data in enumerate(ws_data):
            if i == 0:
                continue
            if ws_data[f'B{i+1}'].value == 0:
                chapter_desc = ws_data[f'C{i+1}'].value
            l = len(f'{data[0].value}')
            if l == 13:
                sub_headings = [None] * 10
                m_data = None
                n_data = None
                o_data = None
                p_data = None
                hscode_indent = ws_data[f'B{i+1}'].value - 1
                if hscode_indent > 0:
                    count = 9
                    while hscode_indent != 0:
                        temp = i + 1
                        while ws_data[f'B{temp}'].value > hscode_indent:
                            temp -= 1
                        sub_headings[count] = ws_data[f'C{temp}'].value
                        hscode_indent -= 1
                        count -= 1
                if hscode is None:
                    hscode = f'{data[0].value}'
                    group = int(hscode[:4])
                    for j in range(i, 1, -1):
                        if len(f"{ws_data[f'A{j}'].value}") == 10:
                            head = f"{ws_data[f'A{j}'].value}"
                            head_index = ws_data[f'B{j}'].value
                    hscode_index = i
                else:
                    hscode = f'{data[0].value}'
                    group = int(hscode[:4])
                    if ws_data[f'B{i+1}'].value == 0:
                        head = hscode[:10]
                        head_index = 0
                    elif ws_data[f'B{i+1}'].value < head_index:
                        head = hscode[:10]
                        head_index = ws_data[f'B{i+1}'].value
                    else:
                        for j in range(i, hscode_index, -1):
                            if ws_data[f'B{j}'].value == 0:
                                head = hscode[:10]
                                head_index = 0
                                break
                            if len(f"{ws_data[f'A{j}'].value}") != 13 and ws_data[f'A{j}'].value is not None:
                                head = f"{ws_data[f'A{j}'].value}"
                                head_index = ws_data[f'B{j}'].value
                                if len(head) != 10 or len(head) != 7:
                                    if '.' in head[-2:]:
                                        head += '0'
                                    if '.' in head[:4]:
                                        head = '0' + head
                                break
                    hscode_index = i
                if ws_data[f'E{i+1}'].value is not None or ws_data[f'F{i+1}'].value is not None or \
                   ws_data[f'G{i+1}'].value is not None or ws_data[f'H{i+1}'].value is not None:
                    m_data = ws_data[f'E{i+1}'].value
                    n_data = ws_data[f'F{i+1}'].value
                    o_data = ws_data[f'G{i+1}'].value
                    p_data = ws_data[f'H{i+1}'].value
                else:
                    j = i
                    found = False
                    while not found:
                        if ws_data[f'E{j}'].value is not None or ws_data[f'F{j}'].value is not None or \
                           ws_data[f'G{j}'].value is not None or ws_data[f'H{j}'].value is not None:
                            if ws_data[f'A{j}'].value in hscodes:
                                pass
                            else:
                                m_data = ws_data[f'E{j}'].value
                                n_data = ws_data[f'F{j}'].value
                                o_data = ws_data[f'G{j}'].value
                                p_data = ws_data[f'H{j}'].value
                                found = True
                        j -= 1

                hscode_desc = ws_data[f'C{i+1}'].value
                ws_target[f'A{target_index}'].value = hscode
                ws_target[f'B{target_index}'].value = group
                ws_target[f'C{target_index}'].value = head
                ws_target[f'D{target_index}'].value = chapter_desc

                while None in sub_headings:
                    sub_headings.remove(None)

                if len(sub_headings) > 7:
                    to_be_joined = sub_headings[6:]
                    temp = ""
                    for st in to_be_joined:
                        temp += f'{st}, '
                    sub_headings[6] = temp[:-2]

                if len(sub_headings) < 7:
                    while len(sub_headings) < 7:
                        sub_headings.insert(0, None)

                for j, k in enumerate(sub_rows):
                    ws_target[f'{k}{target_index}'].value = sub_headings[j]

                ws_target[f'L{target_index}'].value = hscode_desc
                ws_target[f'M{target_index}'].value = m_data
                ws_target[f'N{target_index}'].value = n_data
                ws_target[f'O{target_index}'].value = o_data
                ws_target[f'P{target_index}'].value = p_data
                target_index += 1
                hscodes.append(hscode)

        titles = ['HS Code', 'Group', 'Head', 'Chapter Desc.',
                  'Sub Heading 0', 'Sub Heading 1', 'Sub Heading 2', 'Sub Heading 3', 'Sub Heading 4', 'Sub Heading 5', 'Sub Heading 6',
                  'HS Code Desc.', 'General Rate of Duty', 'Special Rate of Duty', 'Column 2 Rate of Duty', 'Additional Duties']

        target = "ABCDEFGHIJKLMNOP"

        for i, title in zip(target, titles):
            ws_target[f'{i}1'].value = title

        wb_target.save('data.xlsx')
        # wb_hts_data = load_workbook('/home/taraz/Downloads/htsdata.xlsx')
        # hts_data = wb_hts_data.active
        #
        # for i, data in enumerate(hts_data):
        #     vals = {}
        #     desc = [None] * 10
        #     if len(f'{data[0].value}') == 13:
        #         vals['x_hs_code'] = data[0].value
        #         vals['x_hs_code_group'] = data[0].value[:4]
        #         vals['x_hs_code_description'] = data[2].value
        #         if data[1].value == 0:
        #             vals['x_hs_code_head'] = data[0].value[:7]
        #             vals['x_chapter_description'] = data[2].value
        #             raise UserError(f'{vals}')
        #         else:
        #             indent = data[1].value
        #             for j in range(i, -1, -1):
        #                 if hts_data[j][1].value == 0:
        #                     vals['x_chapter_description'] = hts_data[j][2].value
        #                     break
        #                 elif hts_data[j][1].value < indent:
        #                     desc[indent - 2] = hts_data[j][2].value
        #                     indent -= 1
        #
        #             vals['x_heading_description'] = desc[0]
        #             vals['x_level_1_description'] = desc[1]
        #             vals['x_level_2_description'] = desc[2]
        #             vals['x_level_3_description'] = desc[3]
        #             vals['x_level_4_description'] = desc[4]
        #             vals['x_level_5_description'] = desc[5]
        #             vals['x_level_6_description'] = '%s, %s, %s, %s' % (desc[6], desc[7], desc[8], desc[9])
        #
        #         raise UserError(f'{vals}')

    def unlink(self):
        self.x_country_tariff_ids.unlink()
        return super(HsCodeDatabase, self).unlink()


class CountryTariff(models.Model):
    _name = "country.tariff"
    _description = "Country Tariff"
    _rec_name = "x_import_zone_id"

    x_hs_code_id = fields.Many2one(comodel_name='hs.code.database', string='HS Code ID', required=False)

    x_import_zone_id = fields.Many2one(comodel_name='country.import.zones', string='Zone', required=True, tracking=True)
    x_scenario_id = fields.Many2one(comodel_name='country.scenario.rule', string='Scenario',
                                    required=True, tracking=True)

    x_assessed_value = fields.Float(string='AV(%)', compute="_calculate_assessed_value", store=True, tracking=True)

    x_custom_duty = fields.Float(string='CD(%)', tracking=True)
    x_additional_custom_duty = fields.Float(string='ACD(%)', tracking=True)
    x_regulatory_duty = fields.Float(string='RD(%)', tracking=True)
    x_flat_cd = fields.Float(string='Flat CD', tracking=True)

    x_sale_tax = fields.Float(string='ST(%)', tracking=True)
    x_additional_sale_tax = fields.Float(string='AST(%)', tracking=True)
    x_income_tax = fields.Float(string='IT(%)', tracking=True)
    x_cess = fields.Float(string='CESS(%)', tracking=True)
    x_reduced_sale_tax = fields.Float(string='RST(%)', tracking=True)

    x_state = fields.Selection(selection=[
        ('active', 'Active'),
        ('inactive', 'Inactive'), ], default='active', string='State', tracking=True)

    @api.depends('x_scenario_id', 'x_scenario_id.x_landed_cost', 'x_scenario_id.x_insurance')
    def _calculate_assessed_value(self):
        for rec in self:
            rec.x_assessed_value = 100 + rec.x_scenario_id.x_landed_cost + rec.x_scenario_id.x_insurance

    @api.depends('x_scenario_id', 'x_scenario_id.x_landed_cost', 'x_scenario_id.x_insurance')
    def _calculate_assessed_value(self):
        for rec in self:
            rec.x_assessed_value = 100 + rec.x_scenario_id.x_landed_cost + rec.x_scenario_id.x_insurance

    def change_state_active(self):
        for record in self:
            record.x_state = "active"

    def change_state_inactive(self):
        for record in self:
            record.x_state = "inactive"


class FinancialChanges(models.Model):
    _name = "financial.changes"
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _description = "Financial Changes"
    _rec_name = "display_name"
    
    x_hs_code_id = fields.Many2one(comodel_name='hs.code.database', string='HS Code ID', required=False)

    x_country_id = fields.Many2one('res.country', string='Country of Origin', required=True, tracking=True)

    x_change_in = fields.Selection(selection=[
        ('custom_duty', 'CD'),
        ('flat_cd', 'Flat CD'), 
        ('additional_custom_duty', 'ACD'),
        ('regulatory_duty', 'RD'),
        ('sale_tax', 'ST'),
        ('additional_sale_tax', 'AST'),
        ('income_tax', 'IT'),
        ('landed_cost', 'Landed Cost'),
        ('insurance', 'Insurance'),
        ('cess', 'CESS'), ], required=True, string='Change In',)
    x_change = fields.Float(string='Change (%)', tracking=True)

    x_start_date = fields.Date(string='Start Date', required=False, tracking=True)
    x_end_date = fields.Date(string='End Date', required=False, tracking=True)

    display_name = fields.Char(compute="_compute_display_name", store=True)

    @api.depends('x_change_in', 'x_change')
    def _compute_display_name(self):
        for rec in self:
            change_in = dict(rec._fields['x_change_in'].selection).get(rec.x_change_in)
            rec.display_name = '%s @ %s%s' % (change_in, rec.x_change, '%')


class HsCodeFinancialChanges(models.Model):
    _name = "hs.code.financial.changes"
    _description = 'HS Code Financial Changes'
    _rec_name = "display_name"

    x_label = fields.Char(string='Label', required=True, tracking=True)
    x_name = fields.Char(string='Name', required=True, tracking=True)
    display_name = fields.Char(compute="_compute_display_name", store=True)

    @api.depends('x_name', 'x_label')
    def _compute_display_name(self):
        for rec in self:
            rec.display_name = '%s (%s)' % (rec.x_name, rec.x_label)


class HsCodesTags(models.Model):
    _name = "hs.code.tag"
    _description = "HS Code Tags"
    _rec_name = "x_name"

    x_name = fields.Char(string='Name', required=False)
        
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    
    

